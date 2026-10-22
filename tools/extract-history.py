"""Reads the dated price snapshots in the 'Revised Commercial Price List' and 'Revised Bulk Price List' sheets
and writes them as client price history for Supabase.

  python tools/extract-history.py "path/to/MGC LPG Pricing Matrix.xlsx"

Needs db/import.private.json (run extract-workbook.py first) so each row matches a client ID.
Output (git-ignored, real customer data): db/history.private.json

Commercial: five blocks, each with a date, acquisition, hauling and buffer, and five product columns per client
  (50 kg, 22 kg regular, 22 kg forklift, 11 kg, 2.7 kg). Prices are per cylinder, VAT inclusive.
Bulk: seven dated columns of price per kg per client, VAT inclusive.
Zero or blank cells are not a price and are skipped.
"""
import datetime, json, os, re, sys
import openpyxl
from openpyxl.utils import column_index_from_string as ci

SRC = sys.argv[1]
ROOT = os.path.join(os.path.dirname(os.path.abspath(__file__)), '..')
imp = json.load(open(os.path.join(ROOT, 'db', 'import.private.json'), encoding='utf-8'))
by_name = {}
for a in imp['accounts']:
    by_name.setdefault(re.sub(r'\s+', ' ', a['name']).strip().upper(), a)

wb = openpyxl.load_workbook(SRC, data_only=True, read_only=True)
com = wb['Revised Commercial Price List']
blk = wb['Revised Bulk Price List']


def grid(ws, rows, cols):
    out = {}
    for row in ws.iter_rows(min_row=1, max_row=rows, max_col=cols):
        for c in row:
            if c.value not in (None, ''):
                out.setdefault(c.row, {})[c.column_letter] = c.value
    return out


C = grid(com, 192, 110)
B = grid(blk, 45, 45)


def val(g, r, col):
    v = g.get(r, {}).get(col)
    if v is None or isinstance(v, bool):
        return None
    try:
        return float(v)
    except (TypeError, ValueError):
        return None


def cents(v):
    return None if v is None else int(round(v * 100))


def iso(v):
    if isinstance(v, (datetime.datetime, datetime.date)):
        return v.date().isoformat() if isinstance(v, datetime.datetime) else v.isoformat()
    return None


issues, rows = [], []

# ---- Commercial: five blocks ----
# block start column (the 50 kg column), its cost column (start+1) and the date
BLOCKS = [
    ('CF', 'CG', 3, None),     # date in the cost column, row 3
    ('CL', 'CM', 3, None),
    ('CR', 'CS', 3, None),
    ('CX', 'CY', None, '2026-07-22'),  # date written in the label "Acquisition Cost 7/22/2026"
    ('DD', 'DE', 3, None),
]
PRODUCT_OFFSET = [0, 1, 2, 3, 4]  # 50 kg, 22 kg regular, 22 kg forklift, 11 kg, 2.7 kg
SKU_FOR = {1: '22KG_A', 2: '22KG_FL_A', 3: '11KG_MGAS', 4: '2.7KG'}

for start, cost_col, date_row, fixed_date in BLOCKS:
    s = ci(start)
    cost_idx = ci(cost_col)
    date = fixed_date or iso(C.get(date_row, {}).get(cost_col)) or str(C.get(date_row, {}).get(cost_col))
    acq = cents(val(C, 5, cost_col))
    haul = cents(val(C, 6, cost_col))
    buf = cents(val(C, 8, cost_col))
    for r in range(14, 193):
        name = C.get(r, {}).get('B')
        if not name:
            continue
        key = re.sub(r'\s+', ' ', str(name)).strip().upper()
        acc = by_name.get(key)
        if not acc:
            issues.append(f'Commercial {date} row {r} "{name}": no client in the import, skipped')
            continue
        for off in PRODUCT_OFFSET:
            col = openpyxl.utils.get_column_letter(s + off)
            price = val(C, r, col)
            if not price:
                continue
            if off == 0:
                cls = (acc.get('profile', {}).get('contentClass') or 'A').upper()[:1]
                sku = f'50KG_{cls if cls in "ABC" else "A"}'
            else:
                sku = SKU_FOR[off]
            rows.append({'accountId': acc['id'], 'channelId': 'COMMERCIAL', 'effectiveDate': date, 'skuId': sku, 'unit': 'cyl',
                         'grossCentavos': cents(price), 'acqPerKg': acq, 'haulingPerKg': haul, 'bufferPerKg': buf})

# ---- Bulk: seven dated columns, price per kg ----
BULK_COLS = ['AM', 'AN', 'AO', 'AP', 'AQ', 'AR', 'AS']
dates = {}
for col in BULK_COLS:
    text = str(B.get(12, {}).get(col) or '')
    m = re.search(r'(\d{4})-(\d{2})-(\d{2})', text)
    dates[col] = f'{m.group(1)}-{m.group(2)}-{m.group(3)}' if m else text
for r in range(15, 46):
    name = B.get(r, {}).get('C')
    if not name:
        continue
    key = re.sub(r'\s+', ' ', str(name)).strip().upper()
    acc = by_name.get(key)
    if not acc:
        issues.append(f'Bulk row {r} "{name}": no client in the import, skipped')
        continue
    for col in BULK_COLS:
        price = val(B, r, col)
        if not price:
            continue
        rows.append({'accountId': acc['id'], 'channelId': 'BULK', 'effectiveDate': dates[col], 'skuId': 'BULK_KG', 'unit': 'kg',
                     'grossCentavos': cents(price), 'acqPerKg': None, 'haulingPerKg': None, 'bufferPerKg': None})

ids = set(a['id'] for a in imp['accounts'])
out = {'rows': rows}
json.dump(out, open(os.path.join(ROOT, 'db', 'history.private.json'), 'w', encoding='utf-8'), ensure_ascii=False, indent=1)

from collections import Counter
print(f"history rows: {len(rows)}")
print('by channel and date:', dict(Counter((r['channelId'], r['effectiveDate']) for r in rows)))
print('clients with history:', len({r['accountId'] for r in rows}))
print(f'{len(issues)} things to look at:')
for i in issues[:40]:
    print(' -', i)
if len(issues) > 40:
    print(f' ... and {len(issues) - 40} more')
